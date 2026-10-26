
import streamlit as st
from io import BytesIO
from pathlib import Path
from pypdf import PdfReader
from docx import Document
import sqlite3
import pandas as pd
from dlp.audit import DB_PATH
from dlp.pipeline import scan_text
from dlp.hybrid_storage import save_file_hybrid


st.set_page_config(
    page_title="Cloud DLP System",
    page_icon="🛡️",
    layout="wide",
)

st.title("Python-Based Data Loss Prevention System for Cloud")
st.caption("Detect sensitive data and evaluate DLP policies.")

st.info(
    "Upload the file to scan and store it in Azure. "
    "Blob Storage when available, with local storage as a fallback."
)


def extract_text(uploaded_file):
    """Extract readable text from supported file formats."""

    filename = uploaded_file.name.lower()
    file_bytes = uploaded_file.getvalue()

    if filename.endswith(".pdf"):
        reader = PdfReader(BytesIO(file_bytes))
        return "\n".join(
            page.extract_text() or ""
            for page in reader.pages
        )

    if filename.endswith(".docx"):
        document = Document(BytesIO(file_bytes))
        paragraphs = [
            paragraph.text for paragraph in document.paragraphs
        ]

        for table in document.tables:
            for row in table.rows:
                paragraphs.append(
                    " | ".join(cell.text for cell in row.cells)
                )

        return "\n".join(paragraphs)

    if filename.endswith(".doc"):
        raise ValueError(
            "Legacy .doc files are not supported. "
            "Please save the file as .docx and try again."
        )

    return file_bytes.decode("utf-8-sig")


uploaded_file = st.file_uploader(
    "Upload a file to scan",
    type=["txt", "csv", "json", "log", "pdf", "doc", "docx"],
)

if uploaded_file is not None:
    st.write(f"**Selected file:** {uploaded_file.name}")

    if st.button("Scan and store file", type="primary"):
        try:
            extracted_text = extract_text(uploaded_file)

            if not extracted_text.strip():
                st.warning(
                    "No readable text was found. Scanned-image PDFs "
                    "may require OCR."
                )
                st.stop()

            # Scan the extracted text using the existing DLP pipeline.
            result = scan_text(
                extracted_text,
                filename=uploaded_file.name,
            )

            # Store the ORIGINAL file in Azure or use local fallback.
            storage_result = save_file_hybrid(
              file_bytes=uploaded_file.getvalue(),
    filename=uploaded_file.name,
    action=result["action"],
    findings=result["findings"],
    risk=result["risk"],
)

            st.subheader("Scan Results")

            col1, col2, col3 = st.columns(3)

            col1.metric("Risk Level", result["risk"])
            col2.metric("Recommended Action", result["action"])
            col3.metric(
                "Categories Detected",
                len(result["findings"]),
            )

            st.write("### Detection Summary")
            st.json(result["findings"])

            st.write("### Storage Result")
            st.write(
                f"**Storage destination:** "
                f"{storage_result['storage'].upper()}"
            )
            st.write(
                f"**Policy zone:** `{storage_result['destination']}`"
            )
            st.write(
                f"**Fallback used:** "
                f"{storage_result['fallback_used']}"
            )
            st.write(
                f"**Stored location:** `{storage_result['location']}`"
            )

            st.write("### SHA-256 Content Hash")
            st.code(storage_result["file_hash"])

            if storage_result["fallback_used"]:
                st.warning(
                    "Azure storage failed. The file was stored locally "
                    "according to the DLP policy."
                )
            else:
                st.success("File uploaded to Azure successfully.")

            if result["action"] == "BLOCK":
                st.error(
                    "Policy decision: BLOCK. The file is held in "
                    "the blocked storage zone."
                )
            elif result["action"] == "QUARANTINE":
                st.warning(
                    "Policy decision: QUARANTINE. The file is held "
                    "for review."
                )
            elif result["action"] == "FLAG":
                st.warning(
                    "Policy decision: FLAG. The file is held in "
                    "quarantine for review."
                )
            else:
                st.success(
                    "Policy decision: ALLOW. The file is in "
                    "approved storage."
                )

            st.caption(
                  "Files are scanned locally and stored in Azure Blob Storage "
    "when available, with local storage as a fallback. "
    "This system applies DLP policies to uploaded files."
            )

        except UnicodeDecodeError:
            st.error(
                "The text file is not valid UTF-8. "
                "Please convert it to UTF-8 and try again."
            )
        except Exception as exc:
            st.error(f"Scan or storage failed: {exc}")

st.divider()
st.header("Audit Log")

if DB_PATH.exists():
    with sqlite3.connect(DB_PATH) as connection:
        audit_df = pd.read_sql_query(
            """
            SELECT
                timestamp,
                filename,
                risk,
                action,
                destination,
                status
            FROM audit_log
            ORDER BY id DESC
            """,
            connection,
        )

    if audit_df.empty:
        st.info("No audit events recorded yet.")
    else:
        st.dataframe(
            audit_df,
            use_container_width=True,
            hide_index=True,
        )
        st.caption(
            f"Total audit events: {len(audit_df)}"
        )
else:
    st.info("The audit database has not been created yet.")
